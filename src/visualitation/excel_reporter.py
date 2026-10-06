from pathlib import Path

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


class ExcelVisualizer:

    def __init__(self, file_path: str):
        self.file_path = Path(file_path)

    def create_workbook(self) -> None:
        """Create the Excel workbook if it does not exist."""

        self.file_path.parent.mkdir(parents=True, exist_ok=True)

        if not self.file_path.exists():
            workbook = Workbook()
            default_sheet = workbook.active
            default_sheet.title = "Dataset Summary"

            workbook.save(self.file_path)

    def write_dataframe(
        self, data: pd.DataFrame, sheet_name: str, index: bool = True
    ) -> None:
        """Write a DataFrame to an Excel sheet."""

        self.create_workbook()

        workbook = load_workbook(self.file_path)

        if sheet_name in workbook.sheetnames:
            del workbook[sheet_name]

        worksheet = workbook.create_sheet(sheet_name)

        dataframe = data.copy()

        if index:
            dataframe.reset_index(inplace=True)

        for column_index, column_name in enumerate(dataframe.columns, start=1):
            cell = worksheet.cell(row=1, column=column_index, value=column_name)

            cell.font = Font(bold=True)

        for row_index, row in enumerate(dataframe.itertuples(index=False), start=2):
            for column_index, value in enumerate(row, start=1):
                worksheet.cell(row=row_index, column=column_index, value=value)

        self._format_sheet(worksheet)

        workbook.save(self.file_path)

    def write_dict(self, data: dict, sheet_name: str) -> None:
        """Write a dictionary to an Excel sheet."""

        dataframe = pd.DataFrame(data.items(), columns=["Metric", "Value"])

        self.write_dataframe(data=dataframe, sheet_name=sheet_name, index=False)

    def _format_sheet(self, worksheet) -> None:
        """Apply basic formatting to an Excel sheet."""

        header_fill = PatternFill(fill_type="solid", fgColor="D9EAF7")

        for cell in worksheet[1]:
            cell.fill = header_fill
            cell.font = Font(bold=True)

        worksheet.freeze_panes = "A2"

        for column_cells in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column_cells[0].column)

            for cell in column_cells:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))

            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)
